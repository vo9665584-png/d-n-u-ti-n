import os
import json
import requests
from io import BytesIO
from dotenv import load_dotenv
from openai import OpenAI
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Load môi trường
load_dotenv()
client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))

# CONFIG CẤU HÌNH GIAO DIỆN & FONT (Chuẩn Tiếng Việt)
FONT_TITLE = "Segoe UI"
FONT_BODY = "Arial"
COLOR_BG = RGBColor(15, 23, 42)        # Dark Navy (#0F172A)
COLOR_TEXT_MAIN = RGBColor(248, 250, 252)  # Trắng
COLOR_ACCENT = RGBColor(56, 189, 248)      # Xanh Cyan

def get_unsplash_image(query):
    """Tải hình ảnh từ Unsplash"""
    url = f"https://source.unsplash.com/1600x900/?{query.replace(' ', ',')}"
    try:
        res = requests.get(url, timeout=10)
        if res.status_code == 200:
            return BytesIO(res.content)
    except Exception as e:
        print(f"Lỗi tải ảnh: {e}")
    return None

def generate_slide_content(topic, num_slides=5):
    """AI GPT-4o tạo nội dung JSON"""
    print(f"🧠 AI đang tạo nội dung cho: '{topic}'...")
    
    prompt = f"""
    Bạn là chuyên gia Slide. Tạo bài thuyết trình về: '{topic}' gồm {num_slides} slide.
    Trả về định dạng JSON DUY NHẤT có dạng:
    [
      {{
        "slide_num": 1,
        "title": "Tiêu đề ấn tượng",
        "bullets": ["Ý chính 1", "Ý chính 2", "Ý chính 3"],
        "image_keyword": "English keyword for image search"
      }}
    ]
    Nội dung bằng Tiếng Việt chuẩn, ngắn gọn.
    """
    
    response = client.chat.completions.create(
        model="gpt-4o",
        messages=[{"role": "user", "content": prompt}],
    )
    
    content = response.choices[0].message.content
    if "```json" in content:
        content = content.split("```json")[1].split("```")[0]
    return json.loads(content)

def build_powerpoint(slides_data, filename="Output.pptx"):
    """Dựng slide PowerPoint chuẩn 16:9"""
    print("🎨 Đang dựng Slide PowerPoint...")
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    
    for slide_info in slides_data:
        slide = prs.slides.add_slide(blank_layout)
        
        # Nền tối
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = COLOR_BG
        
        # Tiêu đề
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(6.5), Inches(1.2))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = slide_info.get("title", "")
        p.font.name = FONT_TITLE
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT
        
        # Nội dung chữ
        body_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(6.5), Inches(4.5))
        tf_body = body_box.text_frame
        tf_body.word_wrap = True
        
        for i, bullet in enumerate(slide_info.get("bullets", [])):
            p_bullet = tf_body.add_paragraph() if i > 0 else tf_body.paragraphs[0]
            p_bullet.text = f"• {bullet}\n"
            p_bullet.font.name = FONT_BODY
            p_bullet.font.size = Pt(20)
            p_bullet.font.color.rgb = COLOR_TEXT_MAIN
            p_bullet.space_after = Pt(12)

        # Chèn Ảnh bên phải
        img_data = get_unsplash_image(slide_info.get("image_keyword", "technology"))
        if img_data:
            try:
                slide.shapes.add_picture(img_data, Inches(7.8), Inches(1.0), width=Inches(4.8), height=Inches(5.5))
            except Exception as e:
                pass

    prs.save(filename)
    print(f"✅ XONG! File lưu tại: {filename}")

if __name__ == "__main__":
    chud e = input("👉 Nhập chủ đề muốn làm Slide: ")
    so_slide = int(input("👉 Nhập số lượng slide (VD: 5): "))
    
    data = generate_slide_content(chud e, so_slide)
    if isinstance(data, dict):
        data = list(data.values())[0]
        
    build_powerpoint(data, f"{chud e.replace(' ', '_')}.pptx")
